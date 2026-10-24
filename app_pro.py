import gzip
import hashlib
import io
import json
import re
import socket
import zlib
from datetime import date
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlsplit, urlunsplit, quote_plus
from urllib.request import Request, urlopen

import pandas as pd
import streamlit as st

st.set_page_config(page_title="B2B Lead Builder Pro", page_icon="🎯", layout="wide")

# ----------------------------
# Language and interface
# ----------------------------
language = st.sidebar.selectbox("لغة الواجهة / Interface language", ["العربية", "English"], index=0)
AR = language == "العربية"

def tr(ar, en):
    return ar if AR else en

st.markdown(
    f"""<style>
    html, body, [class*="css"] {{ direction: {'rtl' if AR else 'ltr'}; }}
    [data-testid="stAppViewContainer"] {{ direction: {'rtl' if AR else 'ltr'}; }}
    [data-testid="stDataFrame"] {{ direction: ltr; }}
    [data-testid="stMetric"] {{ text-align: {'right' if AR else 'left'}; }}
    </style>""",
    unsafe_allow_html=True,
)

# Stable English field names are used in exports for easy client handoff.
COLUMNS = [
    "Record Type", "Company Name", "Contact Person", "Executive (if listed)", "Job Title",
    "LinkedIn/Profile URL", "Work Email", "Email Status", "Business Phone", "Company Website",
    "Country", "City / Headquarters", "Industry", "Company Size / Employees", "Founded",
    "Source", "Source URL", "Company Search Link", "People Search Link", "Date Collected",
    "Completeness Score", "Completeness Label", "Notes",
]

AR_COLUMNS = {
    "Record Type": "نوع السجل", "Company Name": "اسم الشركة", "Contact Person": "الشخص المسؤول",
    "Executive (if listed)": "المدير التنفيذي (إن توفر)", "Job Title": "المسمى الوظيفي",
    "LinkedIn/Profile URL": "رابط الملف المهني", "Work Email": "البريد المهني",
    "Email Status": "حالة البريد", "Business Phone": "هاتف العمل", "Company Website": "موقع الشركة",
    "Country": "الدولة", "City / Headquarters": "المدينة / المقر", "Industry": "الصناعة",
    "Company Size / Employees": "حجم الشركة / عدد الموظفين", "Founded": "سنة التأسيس",
    "Source": "المصدر", "Source URL": "رابط مصدر البيانات", "Company Search Link": "رابط البحث عن الشركة",
    "People Search Link": "رابط البحث عن الأشخاص", "Date Collected": "تاريخ الجمع",
    "Completeness Score": "درجة اكتمال البيانات", "Completeness Label": "مستوى الاكتمال", "Notes": "ملاحظات",
}

WIKIDATA_API = "https://www.wikidata.org/w/api.php"
WDQS_ENDPOINT = "https://query.wikidata.org/sparql"
USER_AGENT = "B2BLeadBuilderPro/2.0 (public company research app; respectful API usage)"
TODAY = date.today().isoformat()


def empty_df():
    return pd.DataFrame(columns=COLUMNS)


def as_text(value):
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.casefold() in {"nan", "none", "null", "nat"} else text


def normalize_key(value):
    value = as_text(value).casefold()
    return re.sub(r"[^\w]+", " ", value, flags=re.UNICODE).strip()


def normalized_url(value, domain_only=False):
    value = as_text(value)
    if not value:
        return ""
    candidate = value if "://" in value else "https://" + value
    try:
        parts = urlsplit(candidate)
        host = (parts.hostname or "").casefold()
        if host.startswith("www."):
            host = host[4:]
        if domain_only:
            return host
        path = re.sub(r"/+", "/", parts.path or "").rstrip("/")
        return urlunsplit((parts.scheme.casefold() or "https", host, path, "", "")).rstrip("/")
    except Exception:
        return value.casefold().strip().rstrip("/")


def google_url(query):
    return "https://www.google.com/search?q=" + quote_plus(query.strip()) if query.strip() else ""


def company_search_link(company, keyword="", country=""):
    parts = [as_text(company), as_text(keyword), "official website", as_text(country)]
    query = " ".join(part for part in parts if part)
    return google_url(query)


def people_search_link(company, role="", country=""):
    terms = ['site:linkedin.com/in', f'"{as_text(company)}"']
    if as_text(role):
        terms.append(f'"{as_text(role)}"')
    if as_text(country):
        terms.append(as_text(country))
    return google_url(" ".join(term for term in terms if term))


def http_json(url, params=None, data=None, accept="application/json", timeout=30):
    if data is None and params:
        url = url + "?" + urlencode(params)
    body = urlencode(data).encode("utf-8") if data else None
    headers = {
        "User-Agent": USER_AGENT,
        "Accept": accept,
        "Accept-Encoding": "gzip, deflate",
    }
    if body is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded; charset=UTF-8"
    request = Request(url, data=body, headers=headers)
    with urlopen(request, timeout=timeout) as response:
        raw = response.read()
        content_encoding = response.headers.get("Content-Encoding", "").casefold()
        if content_encoding == "gzip":
            raw = gzip.decompress(raw)
        elif content_encoding == "deflate":
            try:
                raw = zlib.decompress(raw)
            except zlib.error:
                raw = zlib.decompress(raw, -zlib.MAX_WBITS)
        return json.loads(raw.decode("utf-8"))


def escape_sparql_literal(value):
    clean = as_text(value).replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ").replace("\r", " ")
    return '"' + clean + '"'


@st.cache_data(ttl=86400, show_spinner=False)
def resolve_country(country_name):
    payload = http_json(
        WIKIDATA_API,
        params={
            "action": "wbsearchentities", "search": country_name, "language": "en",
            "format": "json", "type": "item", "limit": 10,
        },
        timeout=20,
    )
    candidates = payload.get("search", [])
    target = " ".join(country_name.casefold().split())
    for item in candidates:
        label = " ".join(item.get("label", "").casefold().split())
        description = item.get("description", "").casefold()
        if label == target and ("country" in description or "sovereign state" in description):
            return item.get("id"), item.get("label", country_name)
    for item in candidates:
        description = item.get("description", "").casefold()
        if "country" in description or "sovereign state" in description:
            return item.get("id"), item.get("label", country_name)
    return None, None


@st.cache_data(ttl=3600, show_spinner=False)
def fetch_companies(country_qid, keyword, limit, only_with_website):
    if country_qid:
        country_clause = f"?company wdt:P17 wd:{country_qid}.\n  BIND(wd:{country_qid} AS ?country)"
    else:
        country_clause = "OPTIONAL { ?company wdt:P17 ?country. }"

    text_filter = ""
    if as_text(keyword):
        literal = escape_sparql_literal(keyword)
        text_filter = f"""
  FILTER(
    CONTAINS(LCASE(STR(?companyLabel)), LCASE({literal})) ||
    (BOUND(?industryLabel) && CONTAINS(LCASE(STR(?industryLabel)), LCASE({literal})) )
  )"""

    website_filter = "FILTER(BOUND(?website))" if only_with_website else ""
    # The requested result cap limits the number of rows returned, not the total size
    # of the underlying public directory. The public endpoint may time out on broad queries.
    query = f'''SELECT DISTINCT ?company ?companyLabel ?countryLabel ?website ?industryLabel ?hqLabel ?employees ?inception ?executiveLabel WHERE {{
  {country_clause}
  ?company wdt:P31/wdt:P279* ?companyClass.
  VALUES ?companyClass {{ wd:Q4830453 wd:Q783794 }}
  OPTIONAL {{ ?company wdt:P856 ?website. }}
  OPTIONAL {{ ?company wdt:P452 ?industry. }}
  OPTIONAL {{ ?company wdt:P159 ?hq. }}
  OPTIONAL {{ ?company wdt:P1128 ?employees. }}
  OPTIONAL {{ ?company wdt:P571 ?inception. }}
  OPTIONAL {{ ?company wdt:P169 ?executive. }}
  SERVICE wikibase:label {{ bd:serviceParam wikibase:language "en,ar". }}
  {website_filter}
  {text_filter}
}}
LIMIT {int(limit)}'''

    payload = http_json(
        WDQS_ENDPOINT,
        data={"query": query, "format": "json"},
        accept="application/sparql-results+json",
        timeout=58,
    )
    rows = payload.get("results", {}).get("bindings", [])
    records = []
    for row in rows:
        def val(key):
            return as_text(row.get(key, {}).get("value", ""))

        entity_uri = val("company")
        entity_id = entity_uri.rstrip("/").rsplit("/", 1)[-1]
        company_name = val("companyLabel") or entity_id
        country_label = val("countryLabel")
        website = val("website")
        industry = val("industryLabel")
        headquarters = val("hqLabel")
        employees = val("employees")
        founded = val("inception")
        executive = val("executiveLabel")
        source_url = f"https://www.wikidata.org/wiki/{entity_id}" if entity_id else entity_uri
        record = {
            "Record Type": "Company",
            "Company Name": company_name,
            "Contact Person": "",
            "Executive (if listed)": executive,
            "Job Title": "",
            "LinkedIn/Profile URL": "",
            "Work Email": "",
            "Email Status": "Unknown",
            "Business Phone": "",
            "Company Website": website,
            "Country": country_label,
            "City / Headquarters": headquarters,
            "Industry": industry,
            "Company Size / Employees": employees,
            "Founded": founded[:10] if founded else "",
            "Source": "Wikidata",
            "Source URL": source_url,
            "Company Search Link": company_search_link(company_name, keyword, country_label),
            "People Search Link": people_search_link(company_name, "", country_label),
            "Date Collected": TODAY,
            "Completeness Score": "",
            "Completeness Label": "",
            "Notes": "Public record. Confirm details and applicable contact rules before outreach.",
        }
        records.append(record)

    df = pd.DataFrame(records, columns=COLUMNS)
    if df.empty:
        return empty_df()
    # Collapse multiple statements/industry values for the same Wikidata entity.
    df = df.drop_duplicates(subset=["Source URL"], keep="first").reset_index(drop=True)
    return enrich_quality(df)


# Map common column names from the older app, CSVs, and Excel exports into one schema.
ALIASES = {
    "full name": "Contact Person", "contact name": "Contact Person", "person": "Contact Person",
    "company": "Company Name", "company name": "Company Name", "organization": "Company Name",
    "official website": "Company Website", "website": "Company Website", "company website": "Company Website",
    "linkedin url": "LinkedIn/Profile URL", "linkedin/profile url": "LinkedIn/Profile URL", "profile url": "LinkedIn/Profile URL",
    "work email": "Work Email", "email": "Work Email", "email address": "Work Email",
    "business phone": "Business Phone", "phone": "Business Phone", "phone number": "Business Phone",
    "city": "City / Headquarters", "headquarters / city": "City / Headquarters", "headquarters": "City / Headquarters",
    "company size": "Company Size / Employees", "employees": "Company Size / Employees", "employees (if listed)": "Company Size / Employees",
    "date added": "Date Collected", "date collected": "Date Collected", "source url": "Source URL",
    "company source": "Source URL", "people search link": "People Search Link", "company search link": "Company Search Link",
    "ceo / executive (if listed)": "Executive (if listed)", "job title": "Job Title", "email status": "Email Status",
    "industry": "Industry", "country": "Country", "founded": "Founded", "source": "Source", "notes": "Notes",
    "record type": "Record Type", "contact person": "Contact Person", "executive (if listed)": "Executive (if listed)",
    "linkedin/profile url": "LinkedIn/Profile URL", "city / headquarters": "City / Headquarters",
    "company size / employees": "Company Size / Employees", "completeness score": "Completeness Score",
    "completeness label": "Completeness Label", "company website": "Company Website", "date collected": "Date Collected",
}
ALIASES_NORMALIZED = {normalize_key(key): value for key, value in ALIASES.items()}


def normalize_imported(raw_df):
    if raw_df is None or raw_df.empty:
        return empty_df()
    result_rows = []
    for _, source_row in raw_df.fillna("").iterrows():
        record = {col: "" for col in COLUMNS}
        for original_column, raw_value in source_row.items():
            original = as_text(original_column)
            target = original if original in COLUMNS else ALIASES_NORMALIZED.get(normalize_key(original))
            if target in COLUMNS:
                value = as_text(raw_value)
                if value and not record[target]:
                    record[target] = value
        if not any(as_text(record[col]) for col in COLUMNS if col not in {"Completeness Score", "Completeness Label"}):
            continue
        has_contact = any(record[c] for c in ["Contact Person", "Job Title", "LinkedIn/Profile URL", "Work Email"])
        record["Record Type"] = "Contact" if has_contact else (record["Record Type"] or "Company")
        record["Source"] = record["Source"] or "Imported file"
        record["Date Collected"] = record["Date Collected"] or TODAY
        if record["Company Name"] and not record["Company Search Link"]:
            record["Company Search Link"] = company_search_link(record["Company Name"], record["Industry"], record["Country"])
        if record["Company Name"] and not record["People Search Link"]:
            record["People Search Link"] = people_search_link(record["Company Name"], record["Job Title"], record["Country"])
        result_rows.append(record)
    return enrich_quality(pd.DataFrame(result_rows, columns=COLUMNS)) if result_rows else empty_df()


def enrich_quality(df):
    if df is None or df.empty:
        return empty_df()
    df = df.copy().reindex(columns=COLUMNS, fill_value="").fillna("")
    for col in COLUMNS:
        df[col] = df[col].map(as_text)
    for idx, row in df.iterrows():
        if not row["Record Type"]:
            is_contact = any(row[c] for c in ["Contact Person", "Job Title", "LinkedIn/Profile URL", "Work Email"])
            df.at[idx, "Record Type"] = "Contact" if is_contact else "Company"
        if not row["Date Collected"]:
            df.at[idx, "Date Collected"] = TODAY
        if row["Record Type"].casefold() == "contact":
            checks = ["Contact Person", "Job Title", "Company Name", "LinkedIn/Profile URL", "Work Email", "Business Phone", "Company Website", "Country", "Source URL"]
        else:
            checks = ["Company Name", "Country", "Industry", "Company Website", "City / Headquarters", "Company Size / Employees", "Source URL", "Executive (if listed)"]
        score = round(100 * sum(bool(as_text(row[col])) for col in checks) / len(checks))
        df.at[idx, "Completeness Score"] = int(score)
        df.at[idx, "Completeness Label"] = "High completeness" if score >= 75 else ("Medium completeness" if score >= 45 else "Low completeness")
        if not row["Email Status"] and row["Work Email"]:
            df.at[idx, "Email Status"] = "Not verified"
        elif not row["Email Status"]:
            df.at[idx, "Email Status"] = "Unknown"
    return df.reindex(columns=COLUMNS).reset_index(drop=True)


def record_key(row, index):
    is_contact = as_text(row.get("Record Type", "")).casefold() == "contact"
    profile = normalized_url(row.get("LinkedIn/Profile URL", ""))
    email = as_text(row.get("Work Email", "")).casefold()
    company = normalize_key(row.get("Company Name", ""))
    country = normalize_key(row.get("Country", ""))
    person = normalize_key(row.get("Contact Person", ""))
    website = normalized_url(row.get("Company Website", ""), domain_only=True)
    source_url = normalized_url(row.get("Source URL", ""))
    if is_contact:
        if profile:
            return ("profile", profile)
        if email:
            return ("email", email)
        if person and company:
            return ("person-company", person, company, country)
        return ("unique-contact-row", index)
    if source_url and "wikidata.org/wiki/" in source_url:
        return ("wikidata-company", source_url)
    if website:
        return ("company-domain", website)
    if company:
        return ("company-name", company, country)
    return ("unique-row", index)


def merge_and_deduplicate(frames):
    available = [frame for frame in frames if frame is not None and not frame.empty]
    if not available:
        return empty_df(), 0
    combined = pd.concat([frame.reindex(columns=COLUMNS, fill_value="") for frame in available], ignore_index=True).fillna("")
    combined = enrich_quality(combined)
    if combined.empty:
        return empty_df(), 0
    combined["_original_index"] = range(len(combined))
    combined["_score_sort"] = pd.to_numeric(combined["Completeness Score"], errors="coerce").fillna(0)
    combined = combined.sort_values(["_score_sort", "_original_index"], ascending=[False, True], kind="mergesort")
    records = []
    positions = {}
    for _, series in combined.iterrows():
        row = {col: as_text(series.get(col, "")) for col in COLUMNS}
        key = record_key(row, int(series["_original_index"]))
        if key not in positions:
            positions[key] = len(records)
            records.append(row)
            continue
        existing = records[positions[key]]
        for col in COLUMNS:
            if col == "Source":
                parts = [p.strip() for p in (existing[col] + "; " + row[col]).split(";") if p.strip()]
                existing[col] = "; ".join(dict.fromkeys(parts))
            elif not existing[col] and row[col]:
                existing[col] = row[col]
        if existing["Source URL"] and row["Source URL"] and existing["Source URL"] != row["Source URL"]:
            # Keep the primary source link; source labels are combined above.
            pass
    result = enrich_quality(pd.DataFrame(records, columns=COLUMNS))
    return result, max(0, len(combined) - len(result))


def apply_size_filter(df, size_choice):
    if df.empty or size_choice == "Any size / كل الأحجام":
        return df
    values = pd.to_numeric(df["Company Size / Employees"].astype(str).str.replace(",", "", regex=False), errors="coerce")
    ranges = {
        "1–10": (1, 10), "11–50": (11, 50), "51–200": (51, 200),
        "201–1000": (201, 1000), "1000+": (1001, float("inf")),
    }
    if size_choice == "Known employee count / عدد الموظفين معروف":
        return df.loc[values.notna()].copy()
    low, high = ranges.get(size_choice, (float("-inf"), float("inf")))
    return df.loc[values.between(low, high, inclusive="both")].copy()


def make_excel(df):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="B2B Leads")
        ws = writer.sheets["B2B Leads"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163A5F")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for column_cells in ws.columns:
            sample = [len(as_text(cell.value)) for cell in list(column_cells)[:201]]
            max_len = max(sample, default=10)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 36)
        if len(df) > 0:
            table = Table(displayName="B2BLeadsTable", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)

        summary = writer.book.create_sheet("Summary")
        company_count = int((df["Record Type"].astype(str).str.casefold() != "contact").sum()) if not df.empty else 0
        contact_count = int((df["Record Type"].astype(str).str.casefold() == "contact").sum()) if not df.empty else 0
        avg_score = round(pd.to_numeric(df["Completeness Score"], errors="coerce").mean(), 1) if not df.empty else 0
        source_count = df["Source"].astype(str).nunique() if not df.empty else 0
        summary.append(["Metric", "Value"])
        summary.append(["Created on", TODAY])
        summary.append(["Total records", len(df)])
        summary.append(["Company records", company_count])
        summary.append(["Contact records", contact_count])
        summary.append(["Distinct source labels", source_count])
        summary.append(["Average completeness score (%)", avg_score])
        summary.append(["Important", "Completeness is not identity or email verification."])
        for cell in summary[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163A5F")
        summary.column_dimensions["A"].width = 34
        summary.column_dimensions["B"].width = 58
        summary.freeze_panes = "A2"
    return buffer.getvalue()


# ----------------------------
# App layout
# ----------------------------
st.title(tr("🎯 B2B Lead Builder Pro — إنشاء قوائم العملاء", "🎯 B2B Lead Builder Pro"))
st.write(tr(
    "ابحث عن شركات من السجلات العامة، ادمجها مع ملفاتك، أزل التكرارات، وصدّر النتيجة إلى Excel أو CSV.",
    "Find public company records, combine them with your uploaded files, deduplicate, and export to Excel or CSV."
))
st.info(tr(
    "مهم: البحث الآلي هنا يعتمد على سجلات الشركات المتاحة في Wikidata. رابط Google/X-Ray يساعدك في البحث اليدوي ولا تُجمع نتائج Google أو LinkedIn تلقائيًا. البريد والهاتف يظلان فارغين ما لم يقدمهما المصدر أو ملفك. لا يوجد مصدر مجاني يضمن كل الشركات أو أكثر من 1,000 نتيجة صحيحة في كل بحث.",
    "Important: automated discovery currently uses public Wikidata company records. Google/X-Ray links support manual research; Google and LinkedIn results are not scraped. Emails and phones remain blank unless supplied by the source or your file. No free source guarantees every company or 1,000+ valid records in every search."
))

with st.sidebar:
    st.subheader(tr("📂 دمج ملف موجود", "📂 Merge an existing file"))
    uploaded_file = st.file_uploader(
        tr("ارفع Excel أو CSV لإضافته إلى النتائج", "Upload Excel or CSV to merge with search results"),
        type=["csv", "xlsx"],
        help=tr("ستُدمج الصفوف مع نتائج البحث وتُزال التكرارات المعروفة.", "Rows are merged with search results and obvious duplicates are removed."),
    )
    st.caption(tr("الملفات لا تُحفظ في قاعدة بيانات دائمة؛ نزّل نتائجك بعد كل جلسة.", "Files are not stored in a permanent database; download your results after each session."))

# Read uploaded data safely and only reparse it when the file changes.
if uploaded_file is not None:
    uploaded_bytes = uploaded_file.getvalue()
    upload_hash = hashlib.sha256(uploaded_bytes).hexdigest()
    if st.session_state.get("upload_hash") != upload_hash:
        try:
            if uploaded_file.name.casefold().endswith(".csv"):
                imported_raw = pd.read_csv(io.BytesIO(uploaded_bytes), dtype=str, encoding="utf-8-sig").fillna("")
            else:
                imported_raw = pd.read_excel(io.BytesIO(uploaded_bytes), dtype=str).fillna("")
            st.session_state["imported_data"] = normalize_imported(imported_raw)
            st.session_state["upload_hash"] = upload_hash
            st.session_state["upload_error"] = ""
        except Exception:
            st.session_state["imported_data"] = empty_df()
            st.session_state["upload_hash"] = upload_hash
            st.session_state["upload_error"] = tr("تعذر قراءة الملف. تأكد أنه Excel أو CSV سليم.", "Could not read the file. Make sure it is a valid Excel or CSV file.")
else:
    st.session_state["imported_data"] = empty_df()
    st.session_state["upload_hash"] = ""
    st.session_state["upload_error"] = ""

if st.session_state.get("upload_error"):
    st.error(st.session_state["upload_error"])

if "source_results" not in st.session_state:
    st.session_state["source_results"] = empty_df()
if "search_done" not in st.session_state:
    st.session_state["search_done"] = False

with st.form("company_search_form", clear_on_submit=False):
    st.subheader(tr("1. البحث التلقائي عن الشركات", "1. Automatic company search"))
    col1, col2, col3 = st.columns([1, 1, 1])
    with col1:
        country_input = st.text_input(
            tr("الدولة (اتركها فارغة للبحث عالميًا)", "Country (leave blank for worldwide)"),
            value=st.session_state.get("last_country_input", ""),
            placeholder=tr("مثال: Egypt أو Saudi Arabia", "Example: Egypt or Saudi Arabia"),
        )
        keyword = st.text_input(
            tr("المجال أو كلمة البحث", "Industry or keyword"),
            value=st.session_state.get("last_keyword", ""),
            placeholder=tr("مثال: software, marketing, bank", "Example: software, marketing, bank"),
        )
    with col2:
        role = st.text_input(
            tr("المسمى الوظيفي المستهدف (لرابط البحث)", "Target job title (for search links)"),
            value=st.session_state.get("last_role", ""),
            placeholder=tr("مثال: Sales Director, HR Manager", "Example: Sales Director, HR Manager"),
        )
        size_options = [
            "Any size / كل الأحجام", "1–10", "11–50", "51–200", "201–1000", "1000+",
            "Known employee count / عدد الموظفين معروف",
        ]
        size_choice = st.selectbox(
            tr("فلتر حجم الشركة (اختياري)", "Company-size filter (optional)"),
            size_options,
            index=size_options.index(st.session_state.get("last_size_choice", size_options[0])) if st.session_state.get("last_size_choice", size_options[0]) in size_options else 0,
        )
    with col3:
        max_results = st.selectbox(
            tr("الحد الأقصى لنتائج المصدر", "Maximum results from source"),
            [100, 250, 500, 1000, 2000],
            index=2,
            help=tr("العدد هو حد أقصى وليس ضمانًا. البحث العالمي الكبير قد يتجاوز وقت المصدر.", "This is a cap, not a guarantee. Large worldwide searches may time out at the source."),
        )
        only_with_website = st.checkbox(tr("الموقع الرسمي مطلوب", "Require an official website"), value=False)
    search_button = st.form_submit_button(tr("🔎 ابحث وادمج النتائج", "🔎 Search and merge results"), type="primary", use_container_width=True)

if search_button:
    st.session_state["last_country_input"] = country_input.strip()
    st.session_state["last_keyword"] = keyword.strip()
    st.session_state["last_role"] = role.strip()
    st.session_state["last_size_choice"] = size_choice
    try:
        with st.spinner(tr("جاري البحث في قاعدة البيانات العامة…", "Searching the public database…")):
            qid = None
            resolved_country = ""
            if country_input.strip():
                qid, resolved_country = resolve_country(country_input.strip())
                if not qid:
                    raise ValueError(tr(
                        "لم أتعرف على الدولة. اكتب اسمها بالإنجليزية مثل Egypt أو Saudi Arabia.",
                        "Country not recognized. Use an English name such as Egypt or Saudi Arabia."
                    ))
            fetched = fetch_companies(qid, keyword.strip(), int(max_results), bool(only_with_website))
            if not fetched.empty:
                if qid and resolved_country:
                    fetched["Country"] = resolved_country
                for idx, item in fetched.iterrows():
                    country_for_link = as_text(item["Country"]) or resolved_country or country_input.strip()
                    fetched.at[idx, "Company Search Link"] = company_search_link(item["Company Name"], keyword.strip(), country_for_link)
                    fetched.at[idx, "People Search Link"] = people_search_link(item["Company Name"], role.strip(), country_for_link)
            st.session_state["source_results"] = enrich_quality(fetched)
            st.session_state["search_done"] = True
            st.session_state["last_keyword"] = keyword.strip()
            st.success(tr(f"تم جلب {len(fetched):,} سجل من المصدر العام.", f"Retrieved {len(fetched):,} records from the public source."))
    except ValueError as exc:
        st.error(str(exc))
    except HTTPError as exc:
        if exc.code == 429:
            st.error(tr("المصدر يطلب تقليل الطلبات مؤقتًا (429). انتظر قليلًا ثم جرّب بعدد أقل.", "The source is rate-limiting requests (429). Pause, then retry with a smaller result cap."))
        else:
            st.error(tr(f"رفض مصدر البيانات الطلب (HTTP {exc.code}). جرّب عددًا أقل أو أعد المحاولة لاحقًا.", f"The data source rejected the request (HTTP {exc.code}). Try a smaller cap or retry later."))
    except (URLError, TimeoutError, socket.timeout):
        st.error(tr("تعذر الاتصال بالمصدر أو انتهت مهلة البحث. جرّب كلمة أكثر تحديدًا أو حدًا أقل للنتائج.", "Could not connect to the source or the query timed out. Try a more specific keyword or a lower result cap."))
    except Exception:
        st.error(tr("حدث خطأ غير متوقع أثناء البحث. افتح Manage app ثم Logs لعرض التفاصيل، أو أرسل لي صورة الخطأ.", "An unexpected search error occurred. Open Manage app → Logs for details, or send me a screenshot."))

# Combine automatic public-source records with an uploaded file and deduplicate.
source_df = st.session_state.get("source_results", empty_df())
imported_df = st.session_state.get("imported_data", empty_df())
all_df, duplicates_removed = merge_and_deduplicate([imported_df, source_df])
filtered_df = apply_size_filter(all_df, st.session_state.get("last_size_choice", "Any size / كل الأحجام"))

st.divider()
st.subheader(tr("2. النتائج الموحّدة", "2. Combined results"))
metric_cols = st.columns(4)
company_count = int((filtered_df["Record Type"].astype(str).str.casefold() != "contact").sum()) if not filtered_df.empty else 0
contact_count = int((filtered_df["Record Type"].astype(str).str.casefold() == "contact").sum()) if not filtered_df.empty else 0
with metric_cols[0]:
    st.metric(tr("إجمالي السجلات", "Total records"), f"{len(filtered_df):,}")
with metric_cols[1]:
    st.metric(tr("الشركات", "Companies"), f"{company_count:,}")
with metric_cols[2]:
    st.metric(tr("الأشخاص", "Contacts"), f"{contact_count:,}")
with metric_cols[3]:
    st.metric(tr("التكرارات المحذوفة", "Duplicates removed"), f"{duplicates_removed:,}")

if not st.session_state.get("search_done") and imported_df.empty:
    st.info(tr("ابدأ بكتابة كلمة بحث ثم اضغط «ابحث وادمج النتائج»، أو ارفع ملف Excel/CSV من الشريط الجانبي.", "Enter a keyword and click Search and merge, or upload an Excel/CSV file in the sidebar."))
elif filtered_df.empty:
    st.warning(tr("لا توجد نتائج مطابقة حاليًا. جرّب كلمة بحث مختلفة، أو خفّف فلتر حجم الشركة، أو ارفع ملفًا يحتوي على بيانات.", "No matching records right now. Try another keyword, loosen the company-size filter, or upload a data file."))
else:
    avg_score = round(pd.to_numeric(filtered_df["Completeness Score"], errors="coerce").mean(), 1)
    st.caption(tr(
        f"متوسط اكتمال البيانات: {avg_score}% — الدرجة تقيس الحقول المعبأة فقط، ولا تعني أن البيانات مؤكدة.",
        f"Average data completeness: {avg_score}% — the score measures populated fields only, not verification."
    ))
    display_df = filtered_df.copy()
    if AR:
        display_df = display_df.rename(columns=AR_COLUMNS)
        company_website_col = AR_COLUMNS["Company Website"]
        source_url_col = AR_COLUMNS["Source URL"]
        company_search_col = AR_COLUMNS["Company Search Link"]
        people_search_col = AR_COLUMNS["People Search Link"]
        score_col = AR_COLUMNS["Completeness Score"]
    else:
        company_website_col = "Company Website"
        source_url_col = "Source URL"
        company_search_col = "Company Search Link"
        people_search_col = "People Search Link"
        score_col = "Completeness Score"
    st.dataframe(
        display_df,
        use_container_width=True,
        hide_index=True,
        column_config={
            company_website_col: st.column_config.LinkColumn(company_website_col),
            source_url_col: st.column_config.LinkColumn(source_url_col),
            company_search_col: st.column_config.LinkColumn(company_search_col),
            people_search_col: st.column_config.LinkColumn(people_search_col),
            score_col: st.column_config.ProgressColumn(score_col, min_value=0, max_value=100, format="%d%%"),
        },
    )

# Export controls remain present even when there are no rows, so the user can keep a template.
csv_bytes = filtered_df.to_csv(index=False).encode("utf-8-sig")
xlsx_bytes = make_excel(filtered_df)
export1, export2 = st.columns(2)
with export1:
    st.download_button(
        tr("⬇️ تنزيل Excel احترافي", "⬇️ Download formatted Excel"),
        data=xlsx_bytes,
        file_name="b2b_leads_pro.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )
with export2:
    st.download_button(
        tr("⬇️ تنزيل CSV", "⬇️ Download CSV"),
        data=csv_bytes,
        file_name="b2b_leads_pro.csv",
        mime="text/csv",
        use_container_width=True,
    )

with st.expander(tr("مصادر البيانات وحدود النسخة الحالية", "Data sources and current limitations")):
    st.markdown(tr(
        "- **البحث الآلي:** سجلات الشركات العامة المتاحة في [Wikidata](https://www.wikidata.org/).\n- **دمج ملفاتك:** يدعم CSV وExcel، مع تحويل أسماء أعمدة النسخة القديمة إلى تنسيق موحّد.\n- **روابط البحث:** تُنشأ روابط Google/X-Ray لفتح البحث يدويًا فقط؛ لا يتم تنزيل نتائج محركات البحث تلقائيًا.\n- **الجودة:** التكرارات تُكتشف عبر رابط الملف أو البريد، أو هوية الشركة/الموقع. درجة الاكتمال ليست تحققًا من صحة البيانات.\n- **الحدود:** قد تكون بعض الشركات مفقودة أو غير مكتملة، وقد يرفض المصدر الاستعلامات الكبيرة أو يفرض حدود استخدام.",
        "- **Automatic search:** public company records available in [Wikidata](https://www.wikidata.org/).\n- **File merge:** CSV and Excel are supported; columns from the earlier version are mapped into one schema.\n- **Search links:** Google/X-Ray links open manual searches only; search-engine results are not downloaded automatically.\n- **Quality:** duplicates are detected using profile URL or email, or company/site identity. Completeness is not verification.\n- **Limits:** some companies may be missing or incomplete, and the source may reject large queries or enforce usage limits."
    ))
    st.caption(tr("راجع شروط المصادر والقوانين المعمول بها قبل التواصل التجاري أو بيع البيانات.", "Review source terms and applicable laws before commercial outreach or resale."))

st.caption(tr("نصيحة: احتفظ بالنسخة القديمة app_auto.py حتى تتأكد أن النسخة الجديدة تعمل كما تريد.", "Tip: keep the old app_auto.py until you confirm that the new version works as expected."))
