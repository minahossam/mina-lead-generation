import gzip
import hashlib
import io
import json
import re
import socket
import zlib
from datetime import date
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlsplit, urlunsplit, quote_plus
from urllib.request import Request, urlopen

import pandas as pd
import requests
import streamlit as st

st.set_page_config(page_title="B2B Lead Builder Next", page_icon="🎯", layout="wide")

# ----------------------------
# Language and interface
# ----------------------------
language = st.sidebar.selectbox("لغة الواجهة / Interface language", ["العربية", "English"], index=0)
AR = language == "العربية"

def tr(ar, en):
    return ar if AR else en

st.markdown(
    f"""<style>
    html, body, [class*="css"] {{ direction: {'rtl' if AR else 'ltr'}; }}
    [data-testid="stAppViewContainer"] {{ direction: {'rtl' if AR else 'ltr'}; }}
    [data-testid="stDataFrame"] {{ direction: ltr; }}
    [data-testid="stMetric"] {{ text-align: {'right' if AR else 'left'}; }}
    </style>""",
    unsafe_allow_html=True,
)

# Stable English field names are used in exports for easy client handoff.
COLUMNS = [
    "Record Type", "Company Name", "Contact Person", "Executive (if listed)", "Job Title",
    "LinkedIn/Profile URL", "Work Email", "Email Status", "Business Phone", "Company Website",
    "Country", "City / Headquarters", "Industry", "Company Size / Employees", "Founded",
    "Source", "Source URL", "Company Search Link", "People Search Link", "Date Collected",
    "Completeness Score", "Completeness Label", "Notes",
]

AR_COLUMNS = {
    "Record Type": "نوع السجل", "Company Name": "اسم الشركة", "Contact Person": "الشخص المسؤول",
    "Executive (if listed)": "المدير التنفيذي (إن توفر)", "Job Title": "المسمى الوظيفي",
    "LinkedIn/Profile URL": "رابط الملف المهني", "Work Email": "البريد المهني",
    "Email Status": "حالة البريد", "Business Phone": "هاتف العمل", "Company Website": "موقع الشركة",
    "Country": "الدولة", "City / Headquarters": "المدينة / المقر", "Industry": "الصناعة",
    "Company Size / Employees": "حجم الشركة / عدد الموظفين", "Founded": "سنة التأسيس",
    "Source": "المصدر", "Source URL": "رابط مصدر البيانات", "Company Search Link": "رابط البحث عن الشركة",
    "People Search Link": "رابط البحث عن الأشخاص", "Date Collected": "تاريخ الجمع",
    "Completeness Score": "درجة اكتمال البيانات", "Completeness Label": "مستوى الاكتمال", "Notes": "ملاحظات",
}

WIKIDATA_API = "https://www.wikidata.org/w/api.php"
WDQS_ENDPOINT = "https://query.wikidata.org/sparql"
USER_AGENT = "B2BLeadBuilderPro/2.0 (public company research app; respectful API usage)"
TODAY = date.today().isoformat()


def empty_df():
    return pd.DataFrame(columns=COLUMNS)


def as_text(value):
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.casefold() in {"nan", "none", "null", "nat"} else text


def normalize_key(value):
    value = as_text(value).casefold()
    return re.sub(r"[^\w]+", " ", value, flags=re.UNICODE).strip()


def normalized_url(value, domain_only=False):
    value = as_text(value)
    if not value:
        return ""
    candidate = value if "://" in value else "https://" + value
    try:
        parts = urlsplit(candidate)
        host = (parts.hostname or "").casefold()
        if host.startswith("www."):
            host = host[4:]
        if domain_only:
            return host
        path = re.sub(r"/+", "/", parts.path or "").rstrip("/")
        return urlunsplit((parts.scheme.casefold() or "https", host, path, "", "")).rstrip("/")
    except Exception:
        return value.casefold().strip().rstrip("/")


def google_url(query):
    return "https://www.google.com/search?q=" + quote_plus(query.strip()) if query.strip() else ""


def company_search_link(company, keyword="", country=""):
    parts = [as_text(company), as_text(keyword), "official website", as_text(country)]
    query = " ".join(part for part in parts if part)
    return google_url(query)


def people_search_link(company, role="", country=""):
    terms = ["site:linkedin.com/in"]
    if as_text(company):
        terms.append(f'"{as_text(company)}"')
    for role_term in split_values(role):
        terms.append(f'"{role_term}"')
    if as_text(country):
        terms.append(as_text(country))
    return google_url(" ".join(term for term in terms if term))


def http_json(url, params=None, data=None, accept="application/json", timeout=30):
    if data is None and params:
        url = url + "?" + urlencode(params)
    body = urlencode(data).encode("utf-8") if data else None
    headers = {
        "User-Agent": USER_AGENT,
        "Accept": accept,
        "Accept-Encoding": "gzip, deflate",
    }
    if body is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded; charset=UTF-8"
    request = Request(url, data=body, headers=headers)
    with urlopen(request, timeout=timeout) as response:
        raw = response.read()
        content_encoding = response.headers.get("Content-Encoding", "").casefold()
        if content_encoding == "gzip":
            raw = gzip.decompress(raw)
        elif content_encoding == "deflate":
            try:
                raw = zlib.decompress(raw)
            except zlib.error:
                raw = zlib.decompress(raw, -zlib.MAX_WBITS)
        return json.loads(raw.decode("utf-8"))


def escape_sparql_literal(value):
    clean = as_text(value).replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ").replace("\r", " ")
    return '"' + clean + '"'


@st.cache_data(ttl=86400, show_spinner=False)
def resolve_country(country_name):
    payload = http_json(
        WIKIDATA_API,
        params={
            "action": "wbsearchentities", "search": country_name, "language": "en",
            "format": "json", "type": "item", "limit": 10,
        },
        timeout=20,
    )
    candidates = payload.get("search", [])
    target = " ".join(country_name.casefold().split())
    for item in candidates:
        label = " ".join(item.get("label", "").casefold().split())
        description = item.get("description", "").casefold()
        if label == target and ("country" in description or "sovereign state" in description):
            return item.get("id"), item.get("label", country_name)
    for item in candidates:
        description = item.get("description", "").casefold()
        if "country" in description or "sovereign state" in description:
            return item.get("id"), item.get("label", country_name)
    return None, None


@st.cache_data(ttl=3600, show_spinner=False)
def fetch_companies(country_qid, keyword, limit, only_with_website):
    if country_qid:
        country_clause = f"?company wdt:P17 wd:{country_qid}.\n  BIND(wd:{country_qid} AS ?country)"
    else:
        country_clause = "OPTIONAL { ?company wdt:P17 ?country. }"

    text_filter = ""
    if as_text(keyword):
        literal = escape_sparql_literal(keyword)
        text_filter = f"""
  FILTER(
    CONTAINS(LCASE(STR(?companyLabel)), LCASE({literal})) ||
    (BOUND(?industryLabel) && CONTAINS(LCASE(STR(?industryLabel)), LCASE({literal})) )
  )"""

    website_filter = "FILTER(BOUND(?website))" if only_with_website else ""
    # The requested result cap limits the number of rows returned, not the total size
    # of the underlying public directory. The public endpoint may time out on broad queries.
    query = f'''SELECT DISTINCT ?company ?companyLabel ?countryLabel ?website ?industryLabel ?hqLabel ?employees ?inception ?executiveLabel WHERE {{
  {country_clause}
  ?company wdt:P31/wdt:P279* ?companyClass.
  VALUES ?companyClass {{ wd:Q4830453 wd:Q783794 }}
  OPTIONAL {{ ?company wdt:P856 ?website. }}
  OPTIONAL {{ ?company wdt:P452 ?industry. }}
  OPTIONAL {{ ?company wdt:P159 ?hq. }}
  OPTIONAL {{ ?company wdt:P1128 ?employees. }}
  OPTIONAL {{ ?company wdt:P571 ?inception. }}
  OPTIONAL {{ ?company wdt:P169 ?executive. }}
  SERVICE wikibase:label {{ bd:serviceParam wikibase:language "en,ar". }}
  {website_filter}
  {text_filter}
}}
LIMIT {int(limit)}'''

    payload = http_json(
        WDQS_ENDPOINT,
        data={"query": query, "format": "json"},
        accept="application/sparql-results+json",
        timeout=58,
    )
    rows = payload.get("results", {}).get("bindings", [])
    records = []
    for row in rows:
        def val(key):
            return as_text(row.get(key, {}).get("value", ""))

        entity_uri = val("company")
        entity_id = entity_uri.rstrip("/").rsplit("/", 1)[-1]
        company_name = val("companyLabel") or entity_id
        country_label = val("countryLabel")
        website = val("website")
        industry = val("industryLabel")
        headquarters = val("hqLabel")
        employees = val("employees")
        founded = val("inception")
        executive = val("executiveLabel")
        source_url = f"https://www.wikidata.org/wiki/{entity_id}" if entity_id else entity_uri
        record = {
            "Record Type": "Company",
            "Company Name": company_name,
            "Contact Person": "",
            "Executive (if listed)": executive,
            "Job Title": "",
            "LinkedIn/Profile URL": "",
            "Work Email": "",
            "Email Status": "Unknown",
            "Business Phone": "",
            "Company Website": website,
            "Country": country_label,
            "City / Headquarters": headquarters,
            "Industry": industry,
            "Company Size / Employees": employees,
            "Founded": founded[:10] if founded else "",
            "Source": "Wikidata",
            "Source URL": source_url,
            "Company Search Link": company_search_link(company_name, keyword, country_label),
            "People Search Link": people_search_link(company_name, "", country_label),
            "Date Collected": TODAY,
            "Completeness Score": "",
            "Completeness Label": "",
            "Notes": "Public record. Confirm details and applicable contact rules before outreach.",
        }
        records.append(record)

    df = pd.DataFrame(records, columns=COLUMNS)
    if df.empty:
        return empty_df()
    # Collapse multiple statements/industry values for the same Wikidata entity.
    df = df.drop_duplicates(subset=["Source URL"], keep="first").reset_index(drop=True)
    return enrich_quality(df)


def split_values(value):
    """Split a comma/semicolon/newline-delimited field into unique values."""
    parts = re.split(r"[,;\n]+", as_text(value))
    return list(dict.fromkeys(p.strip() for p in parts if p.strip()))


def apollo_person_record(person, keyword="", country_hint="", role_hint=""):
    """Map Apollo People API data into the app's stable export schema."""
    org = person.get("organization") or person.get("current_employer") or {}
    if not isinstance(org, dict):
        org = {}
    full_name = as_text(person.get("name")) or " ".join(
        p for p in [as_text(person.get("first_name")), as_text(person.get("last_name")) or as_text(person.get("last_name_obfuscated"))] if p
    )
    company_name = as_text(org.get("name")) or as_text(person.get("organization_name"))
    profile_url = as_text(person.get("linkedin_url")) or as_text(person.get("profile_url"))
    website = as_text(org.get("website_url")) or as_text(org.get("website")) or as_text(org.get("primary_domain"))
    if website and not website.startswith(("http://", "https://")):
        website = "https://" + website
    email = as_text(person.get("email"))
    if "email_not_unlocked" in email.casefold() or "not_available" in email.casefold():
        email = ""
    email_status = as_text(person.get("email_status")) or ("Returned by API; verify status" if email else "Not returned by search API")
    phone = as_text(person.get("phone_number")) or as_text(person.get("phone"))
    industry = as_text(org.get("industry"))
    if not industry and isinstance(org.get("industries"), list):
        industry = ", ".join(as_text(x) for x in org.get("industries", []) if as_text(x))
    employees = as_text(org.get("estimated_num_employees")) or as_text(org.get("num_employees"))
    org_country = as_text(org.get("country"))
    person_country = as_text(person.get("country"))
    country = org_country or person_country or country_hint
    city = as_text(org.get("city")) or as_text(person.get("city"))
    source_url = profile_url or website or "https://docs.apollo.io/reference/people-api-search"
    return {
        "Record Type": "Contact",
        "Company Name": company_name,
        "Contact Person": full_name,
        "Executive (if listed)": "",
        "Job Title": as_text(person.get("title")) or as_text(person.get("headline")) or role_hint,
        "LinkedIn/Profile URL": profile_url,
        "Work Email": email,
        "Email Status": email_status,
        "Business Phone": phone,
        "Company Website": website,
        "Country": country,
        "City / Headquarters": city,
        "Industry": industry or keyword,
        "Company Size / Employees": employees,
        "Founded": as_text(org.get("founded_year")),
        "Source": "Apollo API",
        "Source URL": source_url,
        "Company Search Link": company_search_link(company_name, keyword, country),
        "People Search Link": google_url(" ".join(
            x for x in ["site:linkedin.com/in", f'"{full_name}"' if full_name else "", f'"{company_name}"' if company_name else "", f'"{role_hint}"' if role_hint else "", country] if x
        )),
        "Date Collected": TODAY,
        "Completeness Score": "",
        "Completeness Label": "",
        "Notes": ("Apollo People Search API. " + ("Surname is partially obfuscated in this search result. " if person.get("last_name_obfuscated") and not person.get("last_name") else "") + ("Apollo indicates email may exist, but search does not return it. " if person.get("has_email") is True and not email else "") + "Search endpoint does not promise email/phone fields; separate permitted enrichment may be needed."),
    }


def fetch_apollo_people(api_key, titles, seniorities, country, keyword, industry_filter, size_choice, limit):
    """Search people through Apollo's official API; do not scrape LinkedIn."""
    if not api_key:
        raise ValueError("Apollo API key is missing. Add your own key in the password field.")
    url = "https://api.apollo.io/api/v1/mixed_people/api_search"
    records = []
    page = 1
    while len(records) < int(limit):
        per_page = min(100, int(limit) - len(records))
        payload = {"page": page, "per_page": per_page}
        if titles:
            payload["person_titles"] = titles
        if seniorities:
            payload["person_seniorities"] = seniorities
        if country:
            payload["organization_locations"] = [country]
        if keyword:
            payload["q_keywords"] = keyword
        industries = split_values(industry_filter)
        if industries:
            # Apollo requires exact industry names from its own taxonomy.
            payload["organization_industries"] = [x.casefold() for x in industries]
        size_ranges = {
            "1–10": ["1,10"], "11–50": ["11,50"], "51–200": ["51,200"],
            "201–1000": ["201,1000"], "1000+": ["1001,5000", "5001,10000", "10001,50000", "50001,100000"],
            "Known employee count / عدد الموظفين معروف": ["1,10", "11,50", "51,200", "201,1000", "1001,5000", "5001,10000", "10001,50000", "50001,100000"],
        }
        if size_choice in size_ranges:
            payload["organization_num_employees_ranges"] = size_ranges[size_choice]
        response = requests.post(
            url,
            json=payload,
            headers={"x-api-key": api_key, "Accept": "application/json", "Content-Type": "application/json"},
            timeout=35,
        )
        if response.status_code >= 400:
            # Avoid showing response bodies that may echo request details or personal information.
            raise RuntimeError(f"Apollo API returned HTTP {response.status_code}. Check API access, endpoint permissions, filters, and your account limits.")
        data = response.json()
        people = data.get("people") or data.get("person_results") or data.get("results") or []
        if not isinstance(people, list) or not people:
            break
        records.extend(apollo_person_record(person, keyword, country, ", ".join(titles)) for person in people if isinstance(person, dict))
        if len(people) < per_page:
            break
        page += 1
    df = pd.DataFrame(records[:int(limit)], columns=COLUMNS) if records else empty_df()
    return enrich_quality(df)


def findymail_find_email(api_key, full_name, domain):
    """Use Findymail's official finder only after explicit user action."""
    response = requests.post(
        "https://app.findymail.com/api/search/name",
        json={"name": full_name, "domain": domain},
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json", "Accept": "application/json"},
        timeout=35,
    )
    if response.status_code >= 400:
        raise RuntimeError(f"Findymail returned HTTP {response.status_code}. Check your key, account balance, and rate limits.")
    data = response.json()
    contact = data.get("contact") or data.get("data", {}).get("contact") or data.get("data") or {}
    if not isinstance(contact, dict):
        contact = {}
    email = as_text(contact.get("email") or data.get("email"))
    if email:
        status = as_text(contact.get("status") or data.get("status"))
        verified = contact.get("verified", data.get("verified", None))
        if verified is True:
            email_status = "Verified by Findymail"
        elif verified is False:
            email_status = "Not verified by Findymail"
        elif status:
            email_status = status
        else:
            email_status = "Returned by Findymail; status not explicit"
    else:
        email_status = "No email returned by Findymail"
    return email, email_status


def findymail_verify_email(api_key, email):
    """Verify an existing email with an authorized Findymail API key."""
    response = requests.post(
        "https://app.findymail.com/api/verify",
        json={"email": email},
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json", "Accept": "application/json"},
        timeout=35,
    )
    if response.status_code >= 400:
        raise RuntimeError(f"Findymail returned HTTP {response.status_code}. Check your key, account balance, and rate limits.")
    data = response.json()
    verified = data.get("verified")
    if verified is True:
        status = "Verified by Findymail"
    elif verified is False:
        status = "Not verified by Findymail"
    else:
        status = as_text(data.get("status")) or "Provider response received; status unclear"
    return status


# Map common column names from the older app, CSVs, and Excel exports into one schema.
ALIASES = {
    "full name": "Contact Person", "contact name": "Contact Person", "person": "Contact Person",
    "company": "Company Name", "company name": "Company Name", "organization": "Company Name",
    "official website": "Company Website", "website": "Company Website", "company website": "Company Website",
    "linkedin url": "LinkedIn/Profile URL", "linkedin/profile url": "LinkedIn/Profile URL", "profile url": "LinkedIn/Profile URL",
    "work email": "Work Email", "email": "Work Email", "email address": "Work Email",
    "business phone": "Business Phone", "phone": "Business Phone", "phone number": "Business Phone",
    "city": "City / Headquarters", "headquarters / city": "City / Headquarters", "headquarters": "City / Headquarters",
    "company size": "Company Size / Employees", "employees": "Company Size / Employees", "employees (if listed)": "Company Size / Employees",
    "date added": "Date Collected", "date collected": "Date Collected", "source url": "Source URL",
    "company source": "Source URL", "people search link": "People Search Link", "company search link": "Company Search Link",
    "ceo / executive (if listed)": "Executive (if listed)", "job title": "Job Title", "email status": "Email Status",
    "industry": "Industry", "country": "Country", "founded": "Founded", "source": "Source", "notes": "Notes",
    "record type": "Record Type", "contact person": "Contact Person", "executive (if listed)": "Executive (if listed)",
    "linkedin/profile url": "LinkedIn/Profile URL", "city / headquarters": "City / Headquarters",
    "company size / employees": "Company Size / Employees", "completeness score": "Completeness Score",
    "completeness label": "Completeness Label", "company website": "Company Website", "date collected": "Date Collected",
}
ALIASES_NORMALIZED = {normalize_key(key): value for key, value in ALIASES.items()}


def normalize_imported(raw_df):
    if raw_df is None or raw_df.empty:
        return empty_df()
    result_rows = []
    for _, source_row in raw_df.fillna("").iterrows():
        record = {col: "" for col in COLUMNS}
        for original_column, raw_value in source_row.items():
            original = as_text(original_column)
            target = original if original in COLUMNS else ALIASES_NORMALIZED.get(normalize_key(original))
            if target in COLUMNS:
                value = as_text(raw_value)
                if value and not record[target]:
                    record[target] = value
        if not any(as_text(record[col]) for col in COLUMNS if col not in {"Completeness Score", "Completeness Label"}):
            continue
        has_contact = any(record[c] for c in ["Contact Person", "Job Title", "LinkedIn/Profile URL", "Work Email"])
        record["Record Type"] = "Contact" if has_contact else (record["Record Type"] or "Company")
        record["Source"] = record["Source"] or "Imported file"
        record["Date Collected"] = record["Date Collected"] or TODAY
        if record["Company Name"] and not record["Company Search Link"]:
            record["Company Search Link"] = company_search_link(record["Company Name"], record["Industry"], record["Country"])
        if record["Company Name"] and not record["People Search Link"]:
            record["People Search Link"] = people_search_link(record["Company Name"], record["Job Title"], record["Country"])
        result_rows.append(record)
    return enrich_quality(pd.DataFrame(result_rows, columns=COLUMNS)) if result_rows else empty_df()


def enrich_quality(df):
    if df is None or df.empty:
        return empty_df()
    df = df.copy().reindex(columns=COLUMNS, fill_value="").fillna("")
    for col in COLUMNS:
        df[col] = df[col].map(as_text)
    for idx, row in df.iterrows():
        if not row["Record Type"]:
            is_contact = any(row[c] for c in ["Contact Person", "Job Title", "LinkedIn/Profile URL", "Work Email"])
            df.at[idx, "Record Type"] = "Contact" if is_contact else "Company"
        if not row["Date Collected"]:
            df.at[idx, "Date Collected"] = TODAY
        if row["Record Type"].casefold() == "contact":
            checks = ["Contact Person", "Job Title", "Company Name", "LinkedIn/Profile URL", "Work Email", "Business Phone", "Company Website", "Country", "Source URL"]
        else:
            checks = ["Company Name", "Country", "Industry", "Company Website", "City / Headquarters", "Company Size / Employees", "Source URL", "Executive (if listed)"]
        score = round(100 * sum(bool(as_text(row[col])) for col in checks) / len(checks))
        df.at[idx, "Completeness Score"] = int(score)
        df.at[idx, "Completeness Label"] = "High completeness" if score >= 75 else ("Medium completeness" if score >= 45 else "Low completeness")
        if not row["Email Status"] and row["Work Email"]:
            df.at[idx, "Email Status"] = "Not verified"
        elif not row["Email Status"]:
            df.at[idx, "Email Status"] = "Unknown"
    return df.reindex(columns=COLUMNS).reset_index(drop=True)


def record_key(row, index):
    is_contact = as_text(row.get("Record Type", "")).casefold() == "contact"
    profile = normalized_url(row.get("LinkedIn/Profile URL", ""))
    email = as_text(row.get("Work Email", "")).casefold()
    company = normalize_key(row.get("Company Name", ""))
    country = normalize_key(row.get("Country", ""))
    person = normalize_key(row.get("Contact Person", ""))
    website = normalized_url(row.get("Company Website", ""), domain_only=True)
    source_url = normalized_url(row.get("Source URL", ""))
    if is_contact:
        if profile:
            return ("profile", profile)
        # Name + company is stable when a missing email is later enriched.
        if person and company:
            return ("person-company", person, company, country)
        if email:
            return ("email", email)
        return ("unique-contact-row", index)
    if source_url and "wikidata.org/wiki/" in source_url:
        return ("wikidata-company", source_url)
    if website:
        return ("company-domain", website)
    if company:
        return ("company-name", company, country)
    return ("unique-row", index)


def merge_and_deduplicate(frames):
    available = [frame for frame in frames if frame is not None and not frame.empty]
    if not available:
        return empty_df(), 0
    combined = pd.concat([frame.reindex(columns=COLUMNS, fill_value="") for frame in available], ignore_index=True).fillna("")
    combined = enrich_quality(combined)
    if combined.empty:
        return empty_df(), 0
    combined["_original_index"] = range(len(combined))
    combined["_score_sort"] = pd.to_numeric(combined["Completeness Score"], errors="coerce").fillna(0)
    combined = combined.sort_values(["_score_sort", "_original_index"], ascending=[False, True], kind="mergesort")
    records = []
    positions = {}
    for _, series in combined.iterrows():
        row = {col: as_text(series.get(col, "")) for col in COLUMNS}
        key = record_key(row, int(series["_original_index"]))
        if key not in positions:
            positions[key] = len(records)
            records.append(row)
            continue
        existing = records[positions[key]]
        for col in COLUMNS:
            if col == "Source":
                parts = [p.strip() for p in (existing[col] + "; " + row[col]).split(";") if p.strip()]
                existing[col] = "; ".join(dict.fromkeys(parts))
            elif not existing[col] and row[col]:
                existing[col] = row[col]
        if existing["Source URL"] and row["Source URL"] and existing["Source URL"] != row["Source URL"]:
            # Keep the primary source link; source labels are combined above.
            pass
    result = enrich_quality(pd.DataFrame(records, columns=COLUMNS))
    return result, max(0, len(combined) - len(result))


def apply_size_filter(df, size_choice):
    if df.empty or size_choice == "Any size / كل الأحجام":
        return df
    values = pd.to_numeric(df["Company Size / Employees"].astype(str).str.replace(",", "", regex=False), errors="coerce")
    ranges = {
        "1–10": (1, 10), "11–50": (11, 50), "51–200": (51, 200),
        "201–1000": (201, 1000), "1000+": (1001, float("inf")),
    }
    if size_choice == "Known employee count / عدد الموظفين معروف":
        return df.loc[values.notna()].copy()
    low, high = ranges.get(size_choice, (float("-inf"), float("inf")))
    return df.loc[values.between(low, high, inclusive="both")].copy()


def make_excel(df):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="B2B Leads")
        ws = writer.sheets["B2B Leads"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163A5F")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for column_cells in ws.columns:
            sample = [len(as_text(cell.value)) for cell in list(column_cells)[:201]]
            max_len = max(sample, default=10)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 36)
        if len(df) > 0:
            table = Table(displayName="B2BLeadsTable", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)

        summary = writer.book.create_sheet("Summary")
        company_count = int((df["Record Type"].astype(str).str.casefold() != "contact").sum()) if not df.empty else 0
        contact_count = int((df["Record Type"].astype(str).str.casefold() == "contact").sum()) if not df.empty else 0
        avg_score = round(pd.to_numeric(df["Completeness Score"], errors="coerce").mean(), 1) if not df.empty else 0
        source_count = df["Source"].astype(str).nunique() if not df.empty else 0
        summary.append(["Metric", "Value"])
        summary.append(["Created on", TODAY])
        summary.append(["Total records", len(df)])
        summary.append(["Company records", company_count])
        summary.append(["Contact records", contact_count])
        summary.append(["Distinct source labels", source_count])
        summary.append(["Average completeness score (%)", avg_score])
        summary.append(["Important", "Completeness is not identity or email verification."])
        for cell in summary[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163A5F")
        summary.column_dimensions["A"].width = 34
        summary.column_dimensions["B"].width = 58
        summary.freeze_panes = "A2"
    return buffer.getvalue()


# ----------------------------
# App layout
# ----------------------------
st.title(tr("🎯 B2B Lead Builder Next — منصة العملاء المحتملين", "🎯 B2B Lead Builder Next"))
st.write(tr(
    "ابحث عن الشركات من مصدر عام مجاني أو عن الأشخاص عبر API رسمي عندما يكون لديك مفتاح صالح، ثم ادمج الملفات ونظّف التكرارات وصدّر تقريرًا جاهزًا.",
    "Search public company records for free or people through an official API when you have a valid key, then merge files, deduplicate, and export a client-ready report."
))
st.info(tr(
    "الوضع المجاني يبحث في سجلات شركات Wikidata. بحث الأشخاص عبر Apollo يتطلب مفتاح API وحسابًا مؤهلًا، ونتائج البحث لا تضمن البريد أو الهاتف. إثراء/التحقق عبر Findymail اختياري وقد يستهلك أرصدة. لا يتم كشط Google أو LinkedIn.",
    "Free mode searches public Wikidata company records. Apollo people search requires an eligible account/API key and does not guarantee email or phone data. Optional Findymail enrichment/verification may consume credits. Google and LinkedIn are not scraped."
))

with st.sidebar:
    st.subheader(tr("🔑 مفاتيح API (اختيارية)", "🔑 Optional API keys"))
    st.caption(tr(
        "اكتب مفتاحك بنفسك في الجلسة فقط. لا تضع المفاتيح داخل GitHub أو داخل الكود. لا تستخدم مفاتيحك على تطبيق عام لا تتحكم في الوصول إليه.",
        "Enter your own keys for this session only. Never put keys in GitHub or source code. Do not use your keys on a public app whose access you do not control."
    ))
    apollo_api_key = st.text_input("Apollo API key", type="password", key="apollo_key_input", help=tr("مفتاحك يُستخدم فقط عند اختيار بحث الأشخاص.", "Used only if you select Apollo people search."))
    findymail_api_key = st.text_input("Findymail API key", type="password", key="findymail_key_input", help=tr("قد يستهلك البحث أو التحقق أرصدة لدى مزود الخدمة.", "Finder/verifier API calls may consume provider credits."))
    st.divider()
    st.subheader(tr("📂 دمج ملف موجود", "📂 Merge an existing file"))
    uploaded_file = st.file_uploader(
        tr("ارفع Excel أو CSV لإضافته إلى النتائج", "Upload Excel or CSV to merge with search results"),
        type=["csv", "xlsx"],
        help=tr("ستُدمج الصفوف مع نتائج البحث وتُزال التكرارات المعروفة.", "Rows will be merged with search results and known duplicates will be removed."),
    )
    st.caption(tr("الملفات لا تُحفظ في قاعدة بيانات دائمة؛ نزّل نتائجك بعد كل جلسة.", "Files are not stored in a permanent database; download results after each session."))

# Read uploaded data safely and only reparse it when the file changes.
if uploaded_file is not None:
    uploaded_bytes = uploaded_file.getvalue()
    upload_hash = hashlib.sha256(uploaded_bytes).hexdigest()
    if st.session_state.get("upload_hash") != upload_hash:
        try:
            if uploaded_file.name.casefold().endswith(".csv"):
                imported_raw = pd.read_csv(io.BytesIO(uploaded_bytes), dtype=str, encoding="utf-8-sig").fillna("")
            else:
                imported_raw = pd.read_excel(io.BytesIO(uploaded_bytes), dtype=str).fillna("")
            st.session_state["imported_data"] = normalize_imported(imported_raw)
            st.session_state["upload_hash"] = upload_hash
            st.session_state["upload_error"] = ""
        except Exception:
            st.session_state["imported_data"] = empty_df()
            st.session_state["upload_hash"] = upload_hash
            st.session_state["upload_error"] = tr("تعذر قراءة الملف. تأكد أنه Excel أو CSV سليم.", "Could not read the file. Make sure it is a valid Excel or CSV file.")
else:
    st.session_state["imported_data"] = empty_df()
    st.session_state["upload_hash"] = ""
    st.session_state["upload_error"] = ""

if st.session_state.get("upload_error"):
    st.error(st.session_state["upload_error"])

if "source_results" not in st.session_state:
    st.session_state["source_results"] = empty_df()
if "search_done" not in st.session_state:
    st.session_state["search_done"] = False
if "provider_results" not in st.session_state:
    st.session_state["provider_results"] = empty_df()

with st.form("lead_search_form", clear_on_submit=False):
    st.subheader(tr("1. بحث تلقائي وتجهيز القائمة", "1. Automated search and list building"))
    search_mode = st.radio(
        tr("مصدر البحث", "Search source"),
        [tr("الشركات — مصدر عام مجاني", "Companies — free public source"), tr("الأشخاص — Apollo API", "People — Apollo API")],
        horizontal=True,
    )
    col1, col2, col3 = st.columns([1, 1, 1])
    with col1:
        country_input = st.text_input(
            tr("الدولة (اتركها فارغة للبحث العالمي)", "Country (leave blank for worldwide)"),
            value=st.session_state.get("last_country_input", ""),
            placeholder=tr("مثال: Egypt أو Saudi Arabia", "Example: Egypt or Saudi Arabia"),
        )
        keyword = st.text_input(
            tr("كلمة مفتاحية", "Keyword"),
            value=st.session_state.get("last_keyword", ""),
            placeholder=tr("مثال: software, marketing, fintech", "Example: software, marketing, fintech"),
        )
        industry_filter = st.text_input(
            tr("الصناعة (فلتر Apollo اختياري)", "Industry (optional Apollo filter)"),
            value=st.session_state.get("last_industry_filter", ""),
            placeholder=tr("اكتب اسم الصناعة الدقيق في Apollo", "Enter exact Apollo industry names"),
            help=tr("في Apollo يجب أن تطابق تسمية الصناعة الرسمية، وإلا اتركها فارغة.", "Apollo requires exact industry taxonomy labels; leave blank if unsure."),
        )
    with col2:
        role = st.text_input(
            tr("المسمى الوظيفي (افصل الأدوار بفاصلة)", "Job title(s), comma-separated"),
            value=st.session_state.get("last_role", ""),
            placeholder=tr("مثال: Sales Director, HR Manager", "Example: Sales Director, HR Manager"),
        )
        seniority_options = ["owner", "founder", "c_suite", "partner", "vp", "head", "director", "manager", "senior", "entry", "intern"]
        seniorities = st.multiselect(
            tr("المستوى الوظيفي (لـ Apollo)", "Seniority (Apollo only)"),
            seniority_options,
            default=st.session_state.get("last_seniorities", []),
            help=tr("اختياري، ويُستخدم فقط في بحث Apollo.", "Optional; used only for Apollo search."),
        )
        size_options = [
            "Any size / كل الأحجام", "1–10", "11–50", "51–200", "201–1000", "1000+",
            "Known employee count / عدد الموظفين معروف",
        ]
        size_choice = st.selectbox(
            tr("حجم الشركة", "Company size"),
            size_options,
            index=size_options.index(st.session_state.get("last_size_choice", size_options[0])) if st.session_state.get("last_size_choice", size_options[0]) in size_options else 0,
        )
    with col3:
        max_results = st.selectbox(
            tr("الحد الأقصى للنتائج", "Maximum results"),
            [100, 250, 500, 1000],
            index=2,
            help=tr("حد أقصى وليس ضمانًا. المصدر المجاني قد يعيد أقل أو تنتهي مهلته.", "A cap, not a guarantee. The free source may return fewer records or time out."),
        )
        only_with_website = st.checkbox(tr("يشترط وجود موقع رسمي (للشركات)", "Require an official website (companies)"), value=False)
    search_button = st.form_submit_button(tr("🚀 ابحث ونظّف النتائج", "🚀 Search and clean results"), type="primary", use_container_width=True)

if search_button:
    st.session_state["last_country_input"] = country_input.strip()
    st.session_state["last_keyword"] = keyword.strip()
    st.session_state["last_industry_filter"] = industry_filter.strip()
    st.session_state["last_role"] = role.strip()
    st.session_state["last_seniorities"] = seniorities
    st.session_state["last_size_choice"] = size_choice
    st.session_state["provider_results"] = empty_df()
    try:
        with st.spinner(tr("جاري البحث وتجهيز القائمة…", "Searching and building your list…")):
            if search_mode == tr("الأشخاص — Apollo API", "People — Apollo API"):
                if not apollo_api_key.strip():
                    raise ValueError(tr(
                        "بحث الأشخاص يحتاج مفتاح Apollo API صالحًا. أضف مفتاحك في الشريط الجانبي أو اختر بحث الشركات المجاني.",
                        "People search requires your valid Apollo API key. Enter it in the sidebar or choose the free company search."
                    ))
                fetched = fetch_apollo_people(
                    apollo_api_key.strip(), split_values(role), seniorities, country_input.strip(),
                    keyword.strip(), industry_filter.strip(), size_choice, int(max_results)
                )
                st.session_state["source_results"] = fetched
                st.session_state["search_done"] = True
                st.success(tr(f"تم جلب {len(fetched):,} سجل أشخاص من Apollo API.", f"Retrieved {len(fetched):,} people records from Apollo API."))
            else:
                qid = None
                resolved_country = ""
                if country_input.strip():
                    qid, resolved_country = resolve_country(country_input.strip())
                    if not qid:
                        raise ValueError(tr(
                            "لم أتعرف على الدولة. اكتب اسمها بالإنجليزية مثل Egypt أو Saudi Arabia.",
                            "Country not recognized. Use an English name such as Egypt or Saudi Arabia."
                        ))
                fetched = fetch_companies(qid, keyword.strip(), int(max_results), bool(only_with_website))
                if not fetched.empty:
                    if qid and resolved_country:
                        fetched["Country"] = resolved_country
                    for idx, item in fetched.iterrows():
                        country_for_link = as_text(item["Country"]) or resolved_country or country_input.strip()
                        fetched.at[idx, "Company Search Link"] = company_search_link(item["Company Name"], keyword.strip(), country_for_link)
                        fetched.at[idx, "People Search Link"] = people_search_link(item["Company Name"], role.strip(), country_for_link)
                st.session_state["source_results"] = enrich_quality(fetched)
                st.session_state["search_done"] = True
                st.success(tr(f"تم جلب {len(fetched):,} سجل شركة من المصدر العام.", f"Retrieved {len(fetched):,} company records from the public source."))
    except ValueError as exc:
        st.error(str(exc))
    except requests.Timeout:
        st.error(tr("انتهت مهلة الاتصال بواجهة API. جرّب عددًا أقل أو أعد المحاولة لاحقًا.", "The API request timed out. Try a smaller cap or retry later."))
    except requests.RequestException:
        st.error(tr("تعذر الاتصال بواجهة API. تحقق من الاتصال ومفتاح الخدمة.", "Could not reach the API. Check connectivity and the provider key."))
    except RuntimeError as exc:
        st.error(str(exc))
    except HTTPError as exc:
        if exc.code == 429:
            st.error(tr("المصدر يطلب تقليل الطلبات مؤقتًا (429). جرّب بعدد أقل لاحقًا.", "The source is rate-limiting requests (429). Retry later with a smaller cap."))
        else:
            st.error(tr(f"رفض مصدر البيانات الطلب (HTTP {exc.code}). تحقق من المدخلات أو أعد المحاولة لاحقًا.", f"The data source rejected the request (HTTP {exc.code}). Check inputs or retry later."))
    except (URLError, TimeoutError, socket.timeout):
        st.error(tr("تعذر الاتصال بالمصدر أو انتهت مهلة البحث. جرّب كلمة أكثر تحديدًا أو حدًا أقل للنتائج.", "Could not connect to the source or the query timed out. Try a more specific keyword or a lower result cap."))
    except Exception:
        st.error(tr("حدث خطأ غير متوقع. افتح Manage app ثم Logs أو أرسل لي صورة الخطأ.", "An unexpected error occurred. Open Manage app → Logs or send a screenshot."))

# Combine automatic public-source records with an uploaded file and deduplicate.
source_df = st.session_state.get("source_results", empty_df())
imported_df = st.session_state.get("imported_data", empty_df())
provider_df = st.session_state.get("provider_results", empty_df())
all_df, duplicates_removed = merge_and_deduplicate([imported_df, source_df, provider_df])

with st.expander(tr("✉️ إثراء البريد والتحقق منه عبر Findymail (اختياري وقد يستهلك أرصدة)", "✉️ Find missing emails / verify existing emails with Findymail (optional; may consume credits)")):
    st.write(tr(
        "استخدم هذا القسم فقط إذا كان لديك مفتاح Findymail صالح وتقبل احتمال خصم أرصدة. البحث عن البريد يعمل للسجلات التي تحتوي على اسم الشخص ونطاق الشركة، والتحقق يعمل على الإيميلات الموجودة. لا تشغّل العملية إلا بعد مراجعة العدد.",
        "Use this only if you have a valid Findymail key and accept possible credit charges. Email finding needs a contact name and company domain; verification checks existing emails. Review the request count before running."
    ))
    enrich_mode = st.radio(
        tr("العملية", "Operation"),
        [tr("العثور على الإيميلات الناقصة", "Find missing emails"), tr("التحقق من الإيميلات الموجودة", "Verify existing emails")],
        horizontal=True,
        key="findymail_mode",
    )
    request_limit = st.selectbox(tr("أقصى عدد من الطلبات في هذه المرة", "Maximum requests this run"), [10, 25, 50], index=0, key="findymail_limit")
    accept_credits = st.checkbox(tr("أفهم أن خدمة Findymail قد تخصم أرصدة من حسابي.", "I understand Findymail may charge credits to my account."), key="findymail_credit_ack")
    if st.button(tr("تشغيل العملية على النتائج الحالية", "Run on current results"), type="secondary", disabled=(not findymail_api_key.strip() or not accept_credits)):
        working = all_df.copy()
        if working.empty:
            st.warning(tr("لا توجد سجلات لمعالجتها. ابحث أولًا أو ارفع ملفًا.", "No records to process. Search first or upload a file."))
        else:
            if "provider_results" not in st.session_state:
                st.session_state["provider_results"] = empty_df()
            eligible = []
            for idx, row in working.iterrows():
                is_contact = as_text(row.get("Record Type")).casefold() == "contact" or any(as_text(row.get(c)) for c in ["Contact Person", "Job Title", "LinkedIn/Profile URL"])
                email = as_text(row.get("Work Email"))
                if not is_contact:
                    continue
                if enrich_mode == tr("العثور على الإيميلات الناقصة", "Find missing emails"):
                    domain = normalized_url(row.get("Company Website", ""), domain_only=True)
                    contact_name = as_text(row.get("Contact Person"))
                    has_full_name = len(contact_name.split()) >= 2 and "*" not in contact_name
                    if as_text(row.get("Source")).find("Apollo API") >= 0 and not has_full_name:
                        continue
                    if not email and contact_name and has_full_name and domain:
                        eligible.append((idx, domain))
                elif email:
                    eligible.append((idx, ""))
            eligible = eligible[:int(request_limit)]
            if not eligible:
                st.warning(tr("لم أجد سجلات مناسبة لهذه العملية. تأكد من وجود اسم الشخص ونطاق الشركة أو بريد حالي.", "No eligible rows found. Ensure rows have a person name and company domain, or an existing email."))
            else:
                progress = st.progress(0)
                successful = 0
                errors = 0
                for n, (idx, domain) in enumerate(eligible, start=1):
                    try:
                        if enrich_mode == tr("العثور على الإيميلات الناقصة", "Find missing emails"):
                            found_email, found_status = findymail_find_email(findymail_api_key.strip(), as_text(working.at[idx, "Contact Person"]), domain)
                            if found_email:
                                working.at[idx, "Work Email"] = found_email
                                working.at[idx, "Email Status"] = found_status
                                existing_source = as_text(working.at[idx, "Source"])
                                working.at[idx, "Source"] = "; ".join(x for x in [existing_source, "Findymail"] if x)
                                working.at[idx, "Notes"] = (as_text(working.at[idx, "Notes"]) + " Email returned via Findymail; provider status shown in Email Status.").strip()
                                successful += 1
                        else:
                            working.at[idx, "Email Status"] = findymail_verify_email(findymail_api_key.strip(), as_text(working.at[idx, "Work Email"]))
                            existing_source = as_text(working.at[idx, "Source"])
                            working.at[idx, "Source"] = "; ".join(x for x in [existing_source, "Findymail verifier"] if x)
                            successful += 1
                    except (requests.RequestException, RuntimeError, ValueError):
                        errors += 1
                    progress.progress(n / len(eligible))
                st.session_state["provider_results"] = enrich_quality(working)
                st.success(tr(f"اكتملت العملية. سجلات تمت معالجتها: {len(eligible)}، تحديثات ناجحة: {successful}، أخطاء: {errors}.", f"Finished. Rows processed: {len(eligible)}, successful updates: {successful}, errors: {errors}."))
                st.rerun()

filtered_df = apply_size_filter(all_df, st.session_state.get("last_size_choice", "Any size / كل الأحجام"))

st.divider()
st.subheader(tr("2. النتائج الموحّدة", "2. Combined results"))
metric_cols = st.columns(4)
company_count = int((filtered_df["Record Type"].astype(str).str.casefold() != "contact").sum()) if not filtered_df.empty else 0
contact_count = int((filtered_df["Record Type"].astype(str).str.casefold() == "contact").sum()) if not filtered_df.empty else 0
with metric_cols[0]:
    st.metric(tr("إجمالي السجلات", "Total records"), f"{len(filtered_df):,}")
with metric_cols[1]:
    st.metric(tr("الشركات", "Companies"), f"{company_count:,}")
with metric_cols[2]:
    st.metric(tr("الأشخاص", "Contacts"), f"{contact_count:,}")
with metric_cols[3]:
    st.metric(tr("التكرارات المحذوفة", "Duplicates removed"), f"{duplicates_removed:,}")

if not st.session_state.get("search_done") and imported_df.empty:
    st.info(tr("ابدأ بكتابة كلمة بحث ثم اضغط «ابحث وادمج النتائج»، أو ارفع ملف Excel/CSV من الشريط الجانبي.", "Enter a keyword and click Search and merge, or upload an Excel/CSV file in the sidebar."))
elif filtered_df.empty:
    st.warning(tr("لا توجد نتائج مطابقة حاليًا. جرّب كلمة بحث مختلفة، أو خفّف فلتر حجم الشركة، أو ارفع ملفًا يحتوي على بيانات.", "No matching records right now. Try another keyword, loosen the company-size filter, or upload a data file."))
else:
    avg_score = round(pd.to_numeric(filtered_df["Completeness Score"], errors="coerce").mean(), 1)
    st.caption(tr(
        f"متوسط اكتمال البيانات: {avg_score}% — الدرجة تقيس الحقول المعبأة فقط، ولا تعني أن البيانات مؤكدة.",
        f"Average data completeness: {avg_score}% — the score measures populated fields only, not verification."
    ))
    display_df = filtered_df.copy()
    if AR:
        display_df = display_df.rename(columns=AR_COLUMNS)
        company_website_col = AR_COLUMNS["Company Website"]
        source_url_col = AR_COLUMNS["Source URL"]
        company_search_col = AR_COLUMNS["Company Search Link"]
        people_search_col = AR_COLUMNS["People Search Link"]
        score_col = AR_COLUMNS["Completeness Score"]
    else:
        company_website_col = "Company Website"
        source_url_col = "Source URL"
        company_search_col = "Company Search Link"
        people_search_col = "People Search Link"
        score_col = "Completeness Score"
    st.dataframe(
        display_df,
        use_container_width=True,
        hide_index=True,
        column_config={
            company_website_col: st.column_config.LinkColumn(company_website_col),
            source_url_col: st.column_config.LinkColumn(source_url_col),
            company_search_col: st.column_config.LinkColumn(company_search_col),
            people_search_col: st.column_config.LinkColumn(people_search_col),
            score_col: st.column_config.ProgressColumn(score_col, min_value=0, max_value=100, format="%d%%"),
        },
    )

# Export controls remain present even when there are no rows, so the user can keep a template.
csv_bytes = filtered_df.to_csv(index=False).encode("utf-8-sig")
xlsx_bytes = make_excel(filtered_df)
export1, export2 = st.columns(2)
with export1:
    st.download_button(
        tr("⬇️ تنزيل Excel احترافي", "⬇️ Download formatted Excel"),
        data=xlsx_bytes,
        file_name="b2b_leads_next.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )
with export2:
    st.download_button(
        tr("⬇️ تنزيل CSV", "⬇️ Download CSV"),
        data=csv_bytes,
        file_name="b2b_leads_next.csv",
        mime="text/csv",
        use_container_width=True,
    )

with st.expander(tr("مصادر البيانات وحدود النسخة الحالية", "Data sources and current limitations")):
    st.markdown(tr(
        "- **بحث الشركات المجاني:** سجلات عامة من [Wikidata](https://www.wikidata.org/)، لكنها ليست قاعدة عالمية كاملة.\n- **بحث الأشخاص:** عبر Apollo People API الرسمي عند إدخال مفتاح صالح؛ بعض الأسماء قد تظهر جزئيًا فقط، والبريد والهاتف لا يضمنهما البحث.\n- **Findymail الاختياري:** العثور على البريد أو التحقق منه قد يستهلك أرصدة؛ لا يعمل إلا بعد موافقتك على تشغيله.\n- **ملفاتك:** يدعم Excel وCSV، توحيد الأعمدة، دمج السجلات، وإزالة التكرار المعروف.\n- **Google/X-Ray وLinkedIn:** روابط بحث يدوي فقط، دون كشط للنتائج أو أتمتة حساب LinkedIn.\n- **الجودة:** درجة الاكتمال تقيس الحقول المملوءة فقط، وليست إثباتًا لصحة الهوية أو البريد أو الهاتف.",
        "- **Free company search:** public Wikidata records; this is not a complete global directory.\n- **People search:** official Apollo People API when you enter a valid key; some names may be partially obfuscated, and search does not guarantee email or phone data.\n- **Optional Findymail:** finding or verifying email may consume credits; it only runs after your explicit confirmation.\n- **Your files:** Excel/CSV import, column normalization, merging, and known-duplicate removal.\n- **Google/X-Ray and LinkedIn:** manual search links only; no search scraping or LinkedIn-account automation.\n- **Quality:** completeness measures populated fields only, not confirmed identity/email/phone."
    ))
    st.caption(tr("راجع شروط المصادر والقوانين المعمول بها قبل التواصل التجاري أو إعادة بيع البيانات.", "Review source terms and applicable laws before commercial outreach or resale."))

st.caption(tr("نصيحة: شغّل هذه النسخة الجديدة كتطبيق منفصل أولًا، واحتفظ بالتطبيق الحالي كنسخة احتياطية.", "Tip: deploy this as a separate app first and keep your current app as a backup."))
